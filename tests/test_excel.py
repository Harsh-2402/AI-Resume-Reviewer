from io import BytesIO

from openpyxl import load_workbook

from models.analysis import (
    AchievementAnalysis, AchievementAssessment, CandidateEvaluation, CertificationAnalysis, CertificationAssessment,
    CertificationVerification, EvidenceAnalysis, ProjectAnalysis, ProjectAssessment, TransferableMatch,
)
from models.candidate import Achievement, CandidateProfile, Certification, Education, Project, TechnicalSkills
from models.github import GitHubAnalysis, RepoAnalysis
from models.jd import JobProfile
from models.result import STAGES, CandidateResult
from models.scoring import ComponentScore, ScoreCard
from reports.excel_report import FINAL_RANKING_COLUMNS, build_excel_report, report_filename
from scoring.ranking import rank_candidates

JD = JobProfile(job_title="Backend Intern", required_skills=["Python", "FastAPI", "PostgreSQL"])


def _card(overall):
    comps = [ComponentScore(key=k, label=k.title(), score=overall if k != "github" else None, weight=0.1, explanation=f"{k} reason")
             for k in ("education", "coursework", "technical_skills", "projects", "github", "certifications", "achievements", "jd_alignment", "learning_potential")]
    return ScoreCard(components=comps, overall_score=overall, classification="Strong Internship Candidate" if overall >= 80 else "Good Internship Candidate",
                     recommendation="Recommend", evidence_confidence="High")


def _result(cid, name, overall, status="Completed"):
    return CandidateResult(
        candidate_id=cid, resume_file=f"{name}.pdf", status=status, scorecard=_card(overall),
        stage_status={s: "completed" for s in STAGES}, processing_time=12.3,
        profile=CandidateProfile(candidate_name=name, email=f"{name}@x.com", github_url=f"https://github.com/{name}",
                                 education=[Education(degree="B.S. CS", university="U", coursework=["Databases"])],
                                 technical_skills=TechnicalSkills(programming_languages=["Python"]),
                                 projects=[Project(name="P1", technologies=["Python"], github_url=f"https://github.com/{name}/p1")],
                                 certifications=[Certification(name="AWS CCP", issuer="AWS", credential_url="https://www.credly.com/x")],
                                 achievements=[Achievement(title="Dean's List")]),
        project_analysis=ProjectAnalysis(assessments=[ProjectAssessment(name="P1", depth="Advanced", jd_relevance="Direct")]),
        certification_analysis=CertificationAnalysis(assessments=[CertificationAssessment(name="AWS CCP", cert_type="Professional Certification", jd_relevance="High")]),
        achievement_analysis=AchievementAnalysis(assessments=[AchievementAssessment(title="Dean's List", category="Academic", significance="Moderate")]),
        github_analysis=GitHubAnalysis(status="analyzed", username=name, url=f"https://github.com/{name}", github_score=70.0,
                                       relevant_repos=[RepoAnalysis(name="p1", relevance="Direct", has_tests=True)]),
        evidence_analysis=EvidenceAnalysis(jd_skill_matrix=[TransferableMatch(jd_skill="Python", status="Direct Match", matched_by="Python"),
                                                            TransferableMatch(jd_skill="FastAPI", status="Transferable", matched_by="Flask"),
                                                            TransferableMatch(jd_skill="PostgreSQL", status="Missing")],
                                           certification_verifications=[CertificationVerification(name="AWS CCP", status="Verified")]),
        evaluation=CandidateEvaluation(top_strength="Projects", primary_concern="None", strengths=["s1"], concerns=["c1"],
                                       interview_questions=["Q1?", "Q2?"]),
    )


def test_excel_has_all_sheets_sorted_and_linked():
    results = [_result("CAND-001", "low", 72.0), _result("CAND-002", "high", 88.5), _result("CAND-003", "mid", 80.0),
               _result("CAND-004", "bad", 0, status="Failed")]
    rank_candidates(results)
    data = build_excel_report(results, JD, {"weights": {"projects": 0.3}})
    wb = load_workbook(BytesIO(data))

    expected = ["Final Ranking", "Candidate Details", "Project Analysis", "GitHub Analysis", "Skills Matrix", "Education",
                "Certifications", "Achievements", "Interview Recommendations", "Processing Log"]
    assert wb.sheetnames[:10] == expected

    ws = wb["Final Ranking"]
    assert [c.value for c in ws[1]] == FINAL_RANKING_COLUMNS
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3  # failed candidate excluded from ranking
    assert [r[0] for r in rows] == [1, 2, 3]
    assert [r[1] for r in rows] == ["CAND-002", "CAND-003", "CAND-001"]
    scores = [r[FINAL_RANKING_COLUMNS.index("Overall Score")] for r in rows]
    assert scores == sorted(scores, reverse=True)
    assert rows[0][FINAL_RANKING_COLUMNS.index("GitHub Score")] == "N/A"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")

    details = wb["Candidate Details"]
    gh_col = [c.value for c in details[1]].index("GitHub") + 1
    cell = details.cell(row=2, column=gh_col)
    assert cell.hyperlink is not None and cell.hyperlink.target == "https://github.com/high"
    assert details.max_row == 5  # 4 candidates incl. failed

    matrix = wb["Skills Matrix"]
    header = [c.value for c in matrix[1]]
    assert header[3:] == ["Python", "FastAPI", "PostgreSQL"]
    first = [c.value for c in matrix[2]]
    assert first[3:] == ["Match", "Transferable", "Missing"]

    assert wb["Project Analysis"].max_row == 4 and wb["Certifications"].max_row == 4 and wb["Achievements"].max_row == 4
    log = wb["Processing Log"]
    log_rows = list(log.iter_rows(min_row=2, values_only=True))
    assert len(log_rows) == 4 and any(r[2] == "Failed" for r in log_rows)
    questions = wb["Interview Recommendations"].cell(row=2, column=11).value
    assert "1. Q1?" in questions and "2. Q2?" in questions
    assert report_filename().startswith("intern_candidate_evaluation_") and report_filename().endswith(".xlsx")


def test_excel_with_no_ranked_candidates_still_builds():
    data = build_excel_report([_result("CAND-001", "x", 0, status="Failed")], JD)
    wb = load_workbook(BytesIO(data))
    assert wb["Final Ranking"].max_row == 1 and wb["Processing Log"].max_row == 2
