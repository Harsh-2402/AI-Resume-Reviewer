import io

import pytest
from docx import Document

from services.pdf_service import DocumentExtractionError, extract_document_text, extract_pdf_text, extract_resume_text
from tests.conftest import SAMPLE_RESUME, make_docx, make_pdf


def test_extracts_text_from_pdf():
    text, warnings = extract_pdf_text(make_pdf(SAMPLE_RESUME))
    assert "Jane Doe" in text
    assert "github.com/janedoe" in text
    assert "Data Structures" in text
    assert warnings == []


def test_blank_pdf_yields_warning():
    text, warnings = extract_pdf_text(make_pdf(" "))
    assert text.strip() == ""
    assert any("scanned" in w for w in warnings)


def test_invalid_pdf_raises():
    with pytest.raises(DocumentExtractionError):
        extract_pdf_text(b"%PDF-1.4 this is not really a pdf")


def test_pdf_from_path(tmp_path):
    path = tmp_path / "r.pdf"
    path.write_bytes(make_pdf("Hello from disk"))
    text, _ = extract_pdf_text(str(path))
    assert "Hello from disk" in text


def test_document_text_docx_and_txt():
    doc = Document()
    doc.add_paragraph("Data Engineering Intern")
    doc.add_paragraph("Required: SQL, Python")
    buf = io.BytesIO()
    doc.save(buf)
    assert "Data Engineering Intern" in extract_document_text(buf.getvalue(), "jd.docx")
    assert "plain text jd" in extract_document_text(b"plain text jd", "jd.txt")
    with pytest.raises(DocumentExtractionError):
        extract_document_text(b"x", "jd.xlsx")


def test_resume_text_dispatches_by_extension(tmp_path):
    (tmp_path / "a.docx").write_bytes(make_docx(SAMPLE_RESUME))
    (tmp_path / "b.txt").write_text(SAMPLE_RESUME)
    (tmp_path / "c.pdf").write_bytes(make_pdf(SAMPLE_RESUME))
    (tmp_path / "short.docx").write_bytes(make_docx("Hi"))
    for name in ("a.docx", "b.txt", "c.pdf"):
        text, warnings = extract_resume_text(str(tmp_path / name))
        assert "Jane Doe" in text and "Data Structures" in text and warnings == []
    _, warnings = extract_resume_text(str(tmp_path / "short.docx"))
    assert warnings and "Little" in warnings[0]
    (tmp_path / "bad.docx").write_bytes(b"not a docx")
    with pytest.raises(DocumentExtractionError):
        extract_resume_text(str(tmp_path / "bad.docx"))
