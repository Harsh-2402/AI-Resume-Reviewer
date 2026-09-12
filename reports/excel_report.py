"""Professional multi-sheet Excel report (openpyxl). Sheet 1 is the recruiter summary sorted by score;
the remaining sheets preserve the evidence behind every number."""
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import config
from models.jd import JobProfile
from models.result import STAGE_LABELS, STAGES, CandidateResult

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CLASS_FILL = {
    "Exceptional Internship Candidate": "C6EFCE",
    "Strong Internship Candidate": "D9EAD3",
    "Good Internship Candidate": "FFF2CC",
    "Potential / Review": "FCE4D6",
    "Weak Match": "F8CBAD",
    "Low Match": "F4CCCC",
}
REC_FILL = {
    "Strongly Recommend": "C6EFCE", "Recommend": "D9EAD3", "Consider": "FFF2CC",
    "Needs Review": "FCE4D6", "Do Not Prioritize": "F4CCCC",
}
STATUS_FILL = {"Completed": "D9EAD3", "Completed with Warnings": "FFF2CC", "Failed": "F4CCCC", "Duplicate": "E7E6E6"}
SCORE_COLUMNS = {
    "Overall Score", "Education Score", "Coursework Score", "Technical Skills Score", "Project Score", "GitHub Score",
    "Certification Score", "Achievement Score", "JD Alignment Score", "Learning Potential", "Score", "Repository Quality",
    "Documentation", "Technology Evidence", "Certification Score", "Education Score",
}
LONG_TEXT_HINTS = ("Description", "Comments", "Why", "Strengths", "Concerns", "Areas", "Questions", "Coursework",
                   "Warnings", "Errors", "Notes", "Evidence", "Explanation", "Problem", "Highlights", "Technologies",
                   "Languages", "Academic", "Reason")


def report_filename() -> str:
    return f"{config.REPORT_FILENAME_PREFIX}_{date.today().isoformat()}.xlsx"


def _na(value: Any) -> Any:
    if value is None:
        return "N/A"
    if isinstance(value, float):
        return round(value, 1)
    if isinstance(value, list):
        return "; ".join(str(v) for v in value if str(v).strip()) or "N/A"
    if isinstance(value, str):
        return value if value.strip() else "N/A"
    return value


def _join(items: list[str], sep: str = "\n") -> str:
    return sep.join(f"• {i}" for i in items if str(i).strip()) if items else "N/A"


def _write_sheet(wb: Workbook, title: str, columns: list[str], rows: list[list[Any]], *,
                 fills: dict[str, dict[str, str]] | None = None, first: bool = False) -> Worksheet:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30

    for row in rows:
        ws.append([_na(v) for v in row])

    n_rows, n_cols = len(rows) + 1, len(columns)
    ws.freeze_panes = "A2"
    if n_cols:
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{max(n_rows, 2)}"

    fills = fills or {}
    for col_idx, name in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        max_len = max([len(str(name))] + [len(str(r[col_idx - 1])) if col_idx - 1 < len(r) and r[col_idx - 1] is not None else 0 for r in rows])
        is_long = any(h in name for h in LONG_TEXT_HINTS)
        ws.column_dimensions[letter].width = min(max(10, max_len + 2), 60 if is_long else 32)
        is_score = name in SCORE_COLUMNS
        for row_idx in range(2, n_rows + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=is_long)
            value = cell.value
            if isinstance(value, str) and value.startswith(("http://", "https://")):
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            elif is_score and isinstance(value, (int, float)):
                cell.number_format = "0.0"
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if name in fills and isinstance(value, str) and value in fills[name]:
                cell.fill = PatternFill("solid", fgColor=fills[name][value])
        if is_score and n_rows > 2:
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{n_rows}",
                ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                               mid_type="num", mid_value=60, mid_color="FFEB84",
                               end_type="num", end_value=100, end_color="63BE7B"),
            )
    return ws


# ─── Row builders ────────────────────────────────────────────────────────────
def _ranked(results: list[CandidateResult]) -> list[CandidateResult]:
    return sorted([r for r in results if r.rank is not None], key=lambda r: r.rank)


def _final_ranking_rows(results: list[CandidateResult]) -> list[list[Any]]:
    rows = []
    for r in _ranked(results):
        p, s = r.profile, r.scorecard
        rows.append([
            r.rank, r.candidate_id, r.name, p.email, p.phone, p.location, p.degree, p.major, p.university,
            p.graduation_year, p.gpa or "Not Provided", s.overall_score, s.classification, s.recommendation,
            s.score_of("education"), s.score_of("coursework"), s.score_of("technical_skills"), s.score_of("projects"),
            s.score_of("github"), s.score_of("certifications"), s.score_of("achievements"), s.score_of("jd_alignment"),
            s.score_of("learning_potential"), s.evidence_confidence, r.evaluation.top_strength, r.evaluation.primary_concern,
        ])
    return rows


FINAL_RANKING_COLUMNS = [
    "Rank", "Candidate ID", "Candidate Name", "Email", "Phone", "Location", "Degree", "Major", "University",
    "Graduation Year", "GPA", "Overall Score", "Classification", "Interview Recommendation", "Education Score",
    "Coursework Score", "Technical Skills Score", "Project Score", "GitHub Score", "Certification Score",
    "Achievement Score", "JD Alignment Score", "Learning Potential", "Evidence Confidence", "Top Strength", "Primary Concern",
]

CANDIDATE_DETAILS_COLUMNS = [
    "Candidate ID", "Name", "Email", "Phone", "Location", "LinkedIn", "GitHub", "Portfolio", "Degree", "Major",
    "University", "GPA", "Graduation Year", "Coursework", "Programming Languages", "Frameworks", "Databases", "Cloud",
    "DevOps", "AI/ML", "Data Engineering", "Tools", "Project Count", "Internship Count", "Certification Count",
    "Achievement Count", "Hackathon Count", "Overall Score", "Classification", "Recommendation", "Confidence",
]


def _details_rows(results: list[CandidateResult]) -> list[list[Any]]:
    rows = []
    for r in _ranked(results) + [x for x in results if x.rank is None]:
        p, t, s = r.profile, r.profile.technical_skills, r.scorecard
        rows.append([
            r.candidate_id, r.name, p.email, p.phone, p.location, p.linkedin_url, p.github_url, p.portfolio_url,
            p.degree, p.major, p.university, p.gpa or "Not Provided", p.graduation_year, ", ".join(p.all_coursework()),
            ", ".join(t.programming_languages), ", ".join(t.frameworks + t.frontend + t.backend), ", ".join(t.databases),
            ", ".join(t.cloud), ", ".join(t.devops), ", ".join(t.ai_ml), ", ".join(t.data_engineering + t.data_science),
            ", ".join(t.tools + t.testing + t.version_control), len(p.projects), len(p.internships), len(p.certifications),
            len(p.achievements) + len(p.awards), len(p.all_hackathons()),
            s.overall_score if r.rank else None, s.classification if r.rank else r.status,
            s.recommendation if r.rank else "N/A", s.evidence_confidence if r.rank else "N/A",
        ])
    return rows


PROJECT_COLUMNS = [
    "Candidate ID", "Candidate Name", "Project Name", "Description", "Technologies", "Problem Solved", "Role", "Team Size",
    "JD Relevance", "Technical Depth", "Complexity", "Ownership Evidence", "Implementation Evidence", "GitHub URL",
    "Live URL", "Demo URL", "Project Score", "Comments",
]


def _project_rows(results: list[CandidateResult]) -> list[list[Any]]:
    from scoring.internship_scorer import _DEPTH_BASE, _IMPL_ADJ, _OWN_ADJ, _REL_MULT, _clamp

    rows = []
    for r in _ranked(results):
        by_name = {a.name.strip().lower(): a for a in r.project_analysis.assessments}
        for p in r.profile.projects:
            a = by_name.get(p.name.strip().lower())
            score = None
            if a:
                score = _clamp(_DEPTH_BASE.get(a.depth, 30) * _REL_MULT.get(a.jd_relevance, 0.6)
                               + _OWN_ADJ.get(a.ownership, 0) + _IMPL_ADJ.get(a.implementation_evidence, 0))
            rows.append([
                r.candidate_id, r.name, p.name, p.description, ", ".join(p.technologies), p.problem_solved, p.role,
                p.team_size, a.jd_relevance if a else None, a.depth if a else None, a.complexity if a else None,
                a.ownership if a else None, a.implementation_evidence if a else None, p.github_url, p.live_url, p.demo_url,
                score, (a.comments + (" Highlights: " + "; ".join(a.technical_highlights) if a.technical_highlights else "")) if a else "Not assessed",
            ])
    return rows


GITHUB_COLUMNS = [
    "Candidate ID", "Candidate Name", "GitHub Username", "GitHub URL", "Public Repositories", "Relevant Repositories",
    "Languages", "Recent Activity", "Meaningful Commits", "Repository Quality", "Documentation", "Testing", "CI/CD",
    "Technology Evidence", "GitHub Score", "Confidence", "Status", "Notes",
]


def _github_rows(results: list[CandidateResult]) -> list[list[Any]]:
    rows = []
    for r in _ranked(results):
        g = r.github_analysis
        relevant = [x for x in g.relevant_repos if x.relevance in ("Direct", "Strongly Related")]
        rows.append([
            r.candidate_id, r.name, g.username or "N/A", g.url, g.public_repos if g.status == "analyzed" else None,
            "; ".join(f"{x.name} ({x.relevance})" for x in relevant) if relevant else ("None" if g.status == "analyzed" else None),
            ", ".join(g.languages), g.last_activity[:10] if g.last_activity else None,
            g.meaningful_commits if g.status == "analyzed" else None,
            g.breakdown.repository_quality if g.status == "analyzed" else None,
            g.breakdown.documentation if g.status == "analyzed" else None,
            ("Yes" if any(x.has_tests for x in g.relevant_repos) else "No") if g.status == "analyzed" else None,
            ("Yes" if any(x.has_ci for x in g.relevant_repos) else "No") if g.status == "analyzed" else None,
            ", ".join(g.technology_evidence), g.github_score, g.confidence, g.status.replace("_", " "),
            "; ".join(g.notes) + (" " + g.explanation if g.explanation else ""),
        ])
    return rows


def _skills_matrix(results: list[CandidateResult], jd: JobProfile) -> tuple[list[str], list[list[Any]]]:
    skills = jd.core_skills()
    columns = ["Rank", "Candidate ID", "Candidate", *skills]
    rows = []
    for r in _ranked(results):
        matrix = {m.jd_skill.strip().lower(): m.status for m in r.evidence_analysis.jd_skill_matrix}
        demonstrated = {a.skill.strip().lower() for a in r.skills_analysis.assessments if a.depth == "Demonstrated"}
        covered_by = {m.jd_skill.strip().lower(): m.matched_by.strip().lower() for m in r.evidence_analysis.jd_skill_matrix}
        row: list[Any] = [r.rank, r.candidate_id, r.name]
        for s in skills:
            status = matrix.get(s.strip().lower(), "N/A")
            if status == "Direct Match":
                row.append("Strong Match" if covered_by.get(s.strip().lower(), "") in demonstrated else "Match")
            elif status == "Strongly Related":
                row.append("Related")
            elif status == "Transferable":
                row.append("Transferable")
            elif status == "Missing":
                row.append("Missing")
            else:
                row.append("N/A")
        rows.append(row)
    return columns, rows


EDUCATION_COLUMNS = ["Rank", "Candidate", "Degree", "Major", "University", "Graduation Year", "GPA", "Coursework",
                     "Academic Projects", "Academic Achievements", "Education Score", "Coursework Score", "Comments"]


def _education_rows(results: list[CandidateResult]) -> list[list[Any]]:
    rows = []
    for r in _ranked(results):
        p, e = r.profile, r.education_analysis
        direct = [m.course for m in e.coursework_matches if m.relevance == "Direct"]
        rows.append([
            r.rank, f"{r.candidate_id} · {r.name}", p.degree, p.major, p.university, p.graduation_year, p.gpa or "Not Provided",
            ", ".join(p.all_coursework()), "; ".join(e.academic_projects or [x.name for x in p.projects if x.is_academic]),
            "; ".join(e.academic_achievements + [h for ed in p.education for h in ed.honors]),
            r.scorecard.score_of("education"), r.scorecard.score_of("coursework"),
            (r.scorecard.component("education").explanation if r.scorecard.component("education") else "")
            + (f" Directly relevant coursework: {', '.join(direct)}." if direct else ""),
        ])
    return rows


CERT_COLUMNS = ["Candidate ID", "Candidate Name", "Certification", "Issuer", "Issue Date", "Expiration Date", "Credential ID",
                "Credential URL", "Technology", "Type", "JD Relevance", "Verification Status", "Certification Score", "Comments"]


def _cert_rows(results: list[CandidateResult]) -> list[list[Any]]:
    from scoring.internship_scorer import _CERT_BASE, _CERT_REL, _VERIFY_BONUS, _clamp

    rows = []
    for r in _ranked(results):
        assess = {a.name.strip().lower(): a for a in r.certification_analysis.assessments}
        verify = {v.name.strip().lower(): v for v in r.evidence_analysis.certification_verifications}
        for c in r.profile.certifications:
            a, v = assess.get(c.name.strip().lower()), verify.get(c.name.strip().lower())
            score = _clamp(_CERT_BASE.get(a.cert_type, 55) * _CERT_REL.get(a.jd_relevance, 0.5) + _VERIFY_BONUS.get(v.status if v else "", 0)) if a else None
            rows.append([
                r.candidate_id, r.name, c.name, c.issuer or (a.issuer if a else ""), c.issue_date, c.expiration_date,
                c.credential_id, c.credential_url, c.technology or (a.technology if a else ""), a.cert_type if a else None,
                a.jd_relevance if a else None, v.status if v else "No Verification Information", score,
                ((a.comments if a else "") + (f" ({v.detail})" if v and v.detail else "")).strip(),
            ])
    return rows


ACH_COLUMNS = ["Candidate", "Achievement", "Category", "Rank / Result", "Organization", "Date", "Relevance", "Evidence", "Score", "Comments"]


def _achievement_rows(results: list[CandidateResult]) -> list[list[Any]]:
    from scoring.internship_scorer import _ACH_BASE, _ACH_REL, _clamp

    rows = []
    for r in _ranked(results):
        for a in r.achievement_analysis.assessments:
            rows.append([
                f"{r.candidate_id} · {r.name}", a.title, a.category, a.rank_or_result, a.organization, a.date, a.relevance,
                a.evidence, _clamp(_ACH_BASE.get(a.significance, 60) * _ACH_REL.get(a.relevance, 0.85)),
                f"Significance: {a.significance}. {a.comments}".strip(),
            ])
    return rows


INTERVIEW_COLUMNS = ["Rank", "Candidate", "Recommendation", "Why Interview", "Key Strengths", "Potential Concerns",
                     "Technical Areas to Test", "Project Areas to Discuss", "GitHub Areas to Discuss",
                     "Certification Areas to Discuss", "Suggested Interview Questions"]


def _interview_rows(results: list[CandidateResult]) -> list[list[Any]]:
    rows = []
    for r in _ranked(results):
        e = r.evaluation
        rows.append([
            r.rank, f"{r.candidate_id} · {r.name}", r.scorecard.recommendation, e.why_interview, _join(e.strengths),
            _join(e.concerns), _join(e.technical_areas_to_test), _join(e.project_areas_to_discuss),
            _join(e.github_areas_to_discuss), _join(e.certification_areas_to_discuss),
            "\n".join(f"{i}. {q}" for i, q in enumerate(e.interview_questions, 1)) or "N/A",
        ])
    return rows


LOG_COLUMNS = ["Candidate ID", "Resume File", "Status", *[STAGE_LABELS[s] for s in STAGES], "Warnings", "Errors", "Processing Time (s)"]


def _log_rows(results: list[CandidateResult]) -> list[list[Any]]:
    label = {"completed": "Completed", "warning": "Completed (warning)", "failed": "Failed", "skipped": "Skipped", "not run": "Not run"}
    rows = []
    for r in sorted(results, key=lambda x: x.candidate_id):
        rows.append([
            r.candidate_id, r.resume_file, r.status + (f" (of {r.duplicate_of})" if r.duplicate_of else ""),
            *[label.get(r.stage_status.get(s, "not run"), r.stage_status.get(s, "")) for s in STAGES],
            _join(r.warnings), _join(r.errors), r.processing_time,
        ])
    return rows


def _jd_rows(jd: JobProfile, settings: dict[str, Any]) -> list[list[Any]]:
    weights = {**config.INTERNSHIP_WEIGHTS, **(settings.get("weights") or {})}
    rows = [["Job Title", jd.job_title], ["Department", jd.department], ["Summary", jd.summary],
            ["Required Skills", ", ".join(jd.required_skills)], ["Preferred Skills", ", ".join(jd.preferred_skills)],
            ["Education Requirements", ", ".join(jd.education_requirements)], ["Relevant Coursework", ", ".join(jd.relevant_coursework)],
            ["Project Expectations", "; ".join(jd.project_expectations)], ["Responsibilities", "; ".join(jd.responsibilities)]]
    rows += [[f"Weight · {config.COMPONENT_LABELS[k]}", f"{v:.0%}"] for k, v in weights.items()]
    rows.append(["Generated", date.today().isoformat()])
    return rows


# ─── Entry point ─────────────────────────────────────────────────────────────
def build_excel_report(results: list[CandidateResult], jd: JobProfile, settings: dict[str, Any] | None = None) -> bytes:
    settings = settings or {}
    wb = Workbook()
    _write_sheet(wb, "Final Ranking", FINAL_RANKING_COLUMNS, _final_ranking_rows(results), first=True,
                 fills={"Classification": CLASS_FILL, "Interview Recommendation": REC_FILL})
    _write_sheet(wb, "Candidate Details", CANDIDATE_DETAILS_COLUMNS, _details_rows(results),
                 fills={"Classification": {**CLASS_FILL, **STATUS_FILL}, "Recommendation": REC_FILL})
    _write_sheet(wb, "Project Analysis", PROJECT_COLUMNS, _project_rows(results))
    _write_sheet(wb, "GitHub Analysis", GITHUB_COLUMNS, _github_rows(results))
    cols, rows = _skills_matrix(results, jd)
    matrix_fill = {"Strong Match": "C6EFCE", "Match": "D9EAD3", "Related": "FFF2CC", "Transferable": "FCE4D6", "Missing": "F4CCCC"}
    _write_sheet(wb, "Skills Matrix", cols, rows, fills={c: matrix_fill for c in cols[3:]})
    _write_sheet(wb, "Education", EDUCATION_COLUMNS, _education_rows(results))
    _write_sheet(wb, "Certifications", CERT_COLUMNS, _cert_rows(results),
                 fills={"Verification Status": {"Verified": "C6EFCE", "Evidence Found": "D9EAD3", "Not Verifiable": "FFF2CC"}})
    _write_sheet(wb, "Achievements", ACH_COLUMNS, _achievement_rows(results))
    _write_sheet(wb, "Interview Recommendations", INTERVIEW_COLUMNS, _interview_rows(results), fills={"Recommendation": REC_FILL})
    _write_sheet(wb, "Processing Log", LOG_COLUMNS, _log_rows(results), fills={"Status": STATUS_FILL})
    _write_sheet(wb, "Job Profile", ["Field", "Value"], _jd_rows(jd, settings))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
