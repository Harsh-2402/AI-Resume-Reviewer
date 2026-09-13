from io import BytesIO
from pathlib import Path

from PyPDF2 import PdfReader
from PyPDF2.errors import PdfReadError

import config
from utils.text import clean_text


class DocumentExtractionError(Exception):
    pass


RESUME_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def resume_skip_reason(filename: str) -> str:
    """Empty string if the file type is accepted as a resume, otherwise a human-readable reason."""
    suffix = Path(filename).suffix.lower()
    if suffix in RESUME_EXTENSIONS:
        return ""
    if suffix == ".doc":
        return "legacy .doc format — save as .docx or PDF"
    return f"unsupported file type ({suffix or 'no extension'})"


def extract_pdf_text(source: str | bytes) -> tuple[str, list[str]]:
    """Return (text, warnings). Raises DocumentExtractionError for unreadable files."""
    warnings: list[str] = []
    try:
        reader = PdfReader(BytesIO(source) if isinstance(source, bytes) else source, strict=False)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise DocumentExtractionError("PDF is password protected") from exc
        pages: list[str] = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001
                warnings.append("A page could not be extracted")
    except DocumentExtractionError:
        raise
    except (PdfReadError, OSError, ValueError, TypeError, KeyError, IndexError, AssertionError) as exc:
        raise DocumentExtractionError(f"Invalid or corrupted PDF: {str(exc)[:120]}") from exc

    text = clean_text("\n".join(pages))
    if len(text) < config.MIN_RESUME_CHARS:
        warnings.append(
            "Little or no extractable text — the PDF may be scanned or image-based; analysis will be limited"
        )
    return text, warnings


def extract_docx_text(source: str | bytes) -> str:
    from docx import Document

    try:
        doc = Document(BytesIO(source) if isinstance(source, bytes) else source)
    except Exception as exc:  # noqa: BLE001
        raise DocumentExtractionError(f"Invalid DOCX: {str(exc)[:120]}") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return clean_text("\n".join(parts))


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return clean_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("Could not decode text file")


def extract_resume_text(path: str) -> tuple[str, list[str]]:
    """Return (text, warnings) for a resume file of any supported type."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(path)
    data = Path(path).read_bytes()
    if suffix == ".docx":
        text = extract_docx_text(data)
    elif suffix in (".txt", ".md"):
        text = _decode_text(data)
    else:
        raise DocumentExtractionError(f"Unsupported file type: {suffix}")
    warnings = [] if len(text) >= config.MIN_RESUME_CHARS else ["Little or no extractable text; analysis will be limited"]
    return text, warnings


def extract_document_text(data: bytes, filename: str) -> str:
    """JD/document text from PDF, DOCX or TXT bytes."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text, _ = extract_pdf_text(data)
        return text
    if suffix == ".docx":
        return extract_docx_text(data)
    if suffix in (".txt", ".md", ""):
        return _decode_text(data)
    raise DocumentExtractionError(f"Unsupported file type: {suffix}")
